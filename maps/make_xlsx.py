import csv, sys
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter

REASON = {
    '천안어린이서점': '주소의 "지하 2"가 도로명 번호 앞에 끼어 있음 — 정확한 번지 확인',
    '금성서점': '"지하204-5" 표기 — 도로명 번호(204-5) 확인',
    '대성문고': '도로명 번호 없음(매실로70번길 뒤 번호 누락) — 정확한 번지 확인',
    '양주어린이서점': '개인주거지로 주소 비공개 — 위치 확인 불가',
}
rows = list(csv.DictReader(open(sys.argv[1], encoding='utf-8-sig')))
cols = list(rows[0].keys())
wb = Workbook(); ws = wb.active; ws.title = '서점목록'
ws.append(cols + ['직접확인'])
fill = PatternFill('solid', fgColor='FFF2CC')
for r in rows:
    flag = r['주소유형'] == '확인필요'
    ws.append([r[c] for c in cols] + [REASON.get(r['상호'], '주소 확인 필요') if flag else ''])
    if flag:
        for cell in ws[ws.max_row]:
            cell.fill = fill
for cell in ws[1]:
    cell.font = Font(bold=True)
ws.freeze_panes = 'A2'
ws.auto_filter.ref = ws.dimensions
for i, c in enumerate(cols + ['직접확인'], 1):
    ws.column_dimensions[get_column_letter(i)].width = {'상호':24,'주소':50,'검색용주소':40,'직접확인':50}.get(c, 16)
wb.save(sys.argv[2])
print(sum(1 for r in rows if r['주소유형']=='확인필요'), '확인필요')
